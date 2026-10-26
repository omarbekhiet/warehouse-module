"""
خدمة استيراد/تصدير Excel
"""
import os
import io
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.database import db


class ExcelService:
    """خدمة Excel الشاملة"""
    
    # =========================================
    # قوالب التصدير
    # =========================================
    EXPORT_TEMPLATES = {
        'accounts': {
            'name': 'الحسابات',
            'sql': """
                SELECT code as 'الكود', name_ar as 'الاسم عربي', 
                       name_en as 'الاسم إنجليزي', level as 'المستوى',
                       nature as 'الطبيعة', parent_code as 'الحساب الأب',
                       financial_statement as 'القائمة', 
                       CASE WHEN is_leaf = 1 THEN 'قابل للترحيل' ELSE 'تجميعي' END as 'النوع',
                       CASE WHEN is_active = 1 THEN 'نشط' ELSE 'موقوف' END as 'الحالة'
                FROM accounts
                ORDER BY code
            """,
            'columns': ['الكود', 'الاسم عربي', 'الاسم إنجليزي', 'المستوى', 
                       'الطبيعة', 'الحساب الأب', 'القائمة', 'النوع', 'الحالة'],
        },
        'customers': {
            'name': 'العملاء',
            'sql': """
                SELECT code as 'الكود', name_ar as 'الاسم', phone as 'الهاتف',
                       email as 'البريد', tax_number as 'الرقم الضريبي',
                       credit_limit as 'حد الائتمان', payment_terms as 'مدة السداد',
                       CASE WHEN is_active = 1 THEN 'نشط' ELSE 'موقوف' END as 'الحالة'
                FROM customers
                ORDER BY code
            """,
            'columns': ['الكود', 'الاسم', 'الهاتف', 'البريد', 'الرقم الضريبي', 
                       'حد الائتمان', 'مدة السداد', 'الحالة'],
        },
        'suppliers': {
            'name': 'الموردين',
            'sql': """
                SELECT code as 'الكود', name_ar as 'الاسم', phone as 'الهاتف',
                       email as 'البريد', tax_number as 'الرقم الضريبي',
                       credit_limit as 'حد الائتمان',
                       CASE WHEN is_active = 1 THEN 'نشط' ELSE 'موقوف' END as 'الحالة'
                FROM suppliers
                ORDER BY code
            """,
            'columns': ['الكود', 'الاسم', 'الهاتف', 'البريد', 'الرقم الضريبي', 
                       'حد الائتمان', 'الحالة'],
        },
        'items': {
            'name': 'الأصناف',
            'sql': """
                SELECT code as 'الكود', name_ar as 'الاسم', 
                       barcode as 'الباركود', unit as 'الوحدة',
                       cost_price as 'سعر التكلفة', sell_price as 'سعر البيع',
                       reorder_level as 'حد الطلب'
                FROM items
                ORDER BY code
            """,
            'columns': ['الكود', 'الاسم', 'الباركود', 'الوحدة', 
                       'سعر التكلفة', 'سعر البيع', 'حد الطلب'],
        },
        'cash_boxes': {
            'name': 'الخزائن',
            'sql': """
                SELECT code as 'الكود', name_ar as 'الاسم', 
                       box_type as 'النوع', account_code as 'الحساب',
                       opening_balance as 'الرصيد الافتتاحي',
                       current_balance as 'الرصيد الحالي',
                       currency as 'العملة'
                FROM cash_boxes
                WHERE is_active = 1
                ORDER BY code
            """,
            'columns': ['الكود', 'الاسم', 'النوع', 'الحساب', 
                       'الرصيد الافتتاحي', 'الرصيد الحالي', 'العملة'],
        },
        'banks': {
            'name': 'البنوك',
            'sql': """
                SELECT code as 'الكود', name_ar as 'الاسم', 
                       account_code as 'الحساب', account_number as 'رقم الحساب',
                       iban as 'IBAN', swift_code as 'SWIFT',
                       branch_name as 'الفرع',
                       current_balance as 'الرصيد'
                FROM banks
                WHERE is_active = 1
                ORDER BY code
            """,
            'columns': ['الكود', 'الاسم', 'الحساب', 'رقم الحساب', 'IBAN', 
                       'SWIFT', 'الفرع', 'الرصيد'],
        },
        'journal_entries': {
            'name': 'القيود',
            'sql': """
                SELECT je.entry_no as 'رقم القيد', je.entry_date as 'التاريخ',
                       je.reference_type as 'نوع المرجع', je.reference_no as 'رقم المرجع',
                       je.description as 'البيان',
                       je.total_debit as 'إجمالي مدين', je.total_credit as 'إجمالي دائن',
                       je.status as 'الحالة'
                FROM journal_entries je
                ORDER BY je.entry_date DESC, je.id DESC
                LIMIT 1000
            """,
            'columns': ['رقم القيد', 'التاريخ', 'نوع المرجع', 'رقم المرجع', 
                       'البيان', 'إجمالي مدين', 'إجمالي دائن', 'الحالة'],
        },
        'supplier_payments': {
            'name': 'مدفوعات الموردين',
            'sql': """
                SELECT sp.payment_no as 'رقم الدفعة', sp.payment_date as 'التاريخ',
                       s.name_ar as 'المورد', sp.payment_method as 'طريقة السداد',
                       sp.amount as 'المبلغ', sp.reference as 'المرجع',
                       je.entry_no as 'رقم القيد'
                FROM supplier_payments sp
                LEFT JOIN suppliers s ON sp.supplier_id = s.id
                LEFT JOIN journal_entries je ON sp.journal_entry_id = je.id
                ORDER BY sp.id DESC
                LIMIT 500
            """,
            'columns': ['رقم الدفعة', 'التاريخ', 'المورد', 'طريقة السداد', 
                       'المبلغ', 'المرجع', 'رقم القيد'],
        },
        'customer_receipts': {
            'name': 'تحصيلات العملاء',
            'sql': """
                SELECT cr.receipt_no as 'رقم التحصيل', cr.receipt_date as 'التاريخ',
                       c.name_ar as 'العميل', cr.receipt_method as 'طريقة التحصيل',
                       cr.amount as 'المبلغ', cr.reference as 'المرجع',
                       je.entry_no as 'رقم القيد'
                FROM customer_receipts cr
                LEFT JOIN customers c ON cr.customer_id = c.id
                LEFT JOIN journal_entries je ON cr.journal_entry_id = je.id
                ORDER BY cr.id DESC
                LIMIT 500
            """,
            'columns': ['رقم التحصيل', 'التاريخ', 'العميل', 'طريقة التحصيل', 
                       'المبلغ', 'المرجع', 'رقم القيد'],
        },
        'cash_box_transactions': {
            'name': 'حركات الخزائن',
            'sql': """
                SELECT cbt.txn_date as 'التاريخ', cbt.txn_type as 'النوع',
                       cbt.reference_no as 'المرجع',
                       COALESCE(cb.name_ar, b.name_ar, pc.name_ar) as 'الخزينة',
                       cbt.amount as 'المبلغ', cbt.balance_after as 'الرصيد بعد',
                       cbt.description as 'البيان'
                FROM cash_box_transactions cbt
                LEFT JOIN cash_boxes cb ON cbt.box_type = 'CASH_BOX' AND cbt.box_id = cb.id
                LEFT JOIN banks b ON cbt.box_type = 'BANK' AND cbt.box_id = b.id
                LEFT JOIN payment_cards pc ON cbt.box_type = 'CARD' AND cbt.box_id = pc.id
                ORDER BY cbt.txn_date DESC, cbt.id DESC
                LIMIT 500
            """,
            'columns': ['التاريخ', 'النوع', 'المرجع', 'الخزينة', 
                       'المبلغ', 'الرصيد بعد', 'البيان'],
        },
        'expenses': {
            'name': 'المصروفات',
            'sql': """
                SELECT e.expense_no as 'رقم المصروف', e.expense_date as 'التاريخ',
                       ec.name_ar as 'التصنيف', e.beneficiary as 'المستفيد',
                       cb.name_ar as 'الخزينة', e.amount as 'المبلغ',
                       e.description as 'البيان', e.reference as 'المرجع'
                FROM petty_cash_expenses e
                LEFT JOIN expense_categories ec ON e.category_id = ec.id
                LEFT JOIN cash_boxes cb ON e.cash_box_id = cb.id
                ORDER BY e.expense_date DESC, e.id DESC
                LIMIT 500
            """,
            'columns': ['رقم المصروف', 'التاريخ', 'التصنيف', 'المستفيد', 
                       'الخزينة', 'المبلغ', 'البيان', 'المرجع'],
        },
        'purchase_invoices': {
            'name': 'فواتير الشراء',
            'sql': """
                SELECT pi.invoice_no as 'رقم الفاتورة', pi.invoice_date as 'التاريخ',
                       s.name_ar as 'المورد', w.name_ar as 'المستودع',
                       pi.total_amount as 'الإجمالي', pi.payment_type as 'نوع الدفع',
                       pi.status as 'الحالة'
                FROM purchase_invoices pi
                LEFT JOIN suppliers s ON pi.supplier_id = s.id
                LEFT JOIN warehouses w ON pi.warehouse_id = w.id
                ORDER BY pi.invoice_date DESC, pi.id DESC
                LIMIT 1000
            """,
            'columns': ['رقم الفاتورة', 'التاريخ', 'المورد', 'المستودع', 
                       'الإجمالي', 'نوع الدفع', 'الحالة'],
        },
        'sales_invoices': {
            'name': 'فواتير البيع',
            'sql': """
                SELECT si.invoice_no as 'رقم الفاتورة', si.invoice_date as 'التاريخ',
                       c.name_ar as 'العميل', w.name_ar as 'المستودع',
                       si.total_amount as 'الإجمالي', si.payment_type as 'نوع الدفع',
                       si.status as 'الحالة'
                FROM sales_invoices si
                LEFT JOIN customers c ON si.customer_id = c.id
                LEFT JOIN warehouses w ON si.warehouse_id = w.id
                ORDER BY si.invoice_date DESC, si.id DESC
                LIMIT 1000
            """,
            'columns': ['رقم الفاتورة', 'التاريخ', 'العميل', 'المستودع', 
                       'الإجمالي', 'نوع الدفع', 'الحالة'],
        },
        'purchase_returns': {
            'name': 'مرتجع شراء',
            'sql': """
                SELECT pr.return_no as 'رقم المرتجع', pr.return_date as 'التاريخ',
                       s.name_ar as 'المورد', pr.total_amount as 'الإجمالي',
                       pr.status as 'الحالة'
                FROM purchase_returns pr
                LEFT JOIN suppliers s ON pr.supplier_id = s.id
                ORDER BY pr.return_date DESC, pr.id DESC
                LIMIT 500
            """,
            'columns': ['رقم المرتجع', 'التاريخ', 'المورد', 'الإجمالي', 'الحالة'],
        },
        'sales_returns': {
            'name': 'مرتجع بيع',
            'sql': """
                SELECT sr.return_no as 'رقم المرتجع', sr.return_date as 'التاريخ',
                       c.name_ar as 'العميل', sr.total_amount as 'الإجمالي',
                       sr.status as 'الحالة'
                FROM sales_returns sr
                LEFT JOIN customers c ON sr.customer_id = c.id
                ORDER BY sr.return_date DESC, sr.id DESC
                LIMIT 500
            """,
            'columns': ['رقم المرتجع', 'التاريخ', 'العميل', 'الإجمالي', 'الحالة'],
        },
        'branches': {
            'name': 'الفروع',
            'sql': """
                SELECT code as 'الكود', name_ar as 'الاسم', phone as 'الهاتف',
                       manager_name as 'المدير', address as 'العنوان',
                       CASE WHEN is_active = 1 THEN 'نشط' ELSE 'موقوف' END as 'الحالة'
                FROM branches
                ORDER BY code
            """,
            'columns': ['الكود', 'الاسم', 'الهاتف', 'المدير', 'العنوان', 'الحالة'],
        },
        'expense_categories': {
            'name': 'تصنيفات المصروفات',
            'sql': """
                SELECT ec.code as 'الكود', ec.name_ar as 'الاسم',
                       ec.account_code as 'الحساب', a.name_ar as 'اسم الحساب',
                       CASE WHEN ec.is_active = 1 THEN 'نشط' ELSE 'موقوف' END as 'الحالة'
                FROM expense_categories ec
                LEFT JOIN accounts a ON ec.account_code = a.code
                ORDER BY ec.code
            """,
            'columns': ['الكود', 'الاسم', 'الحساب', 'اسم الحساب', 'الحالة'],
        },
        'stock_transactions': {
            'name': 'حركات المخزون',
            'sql': """
                SELECT st.txn_date as 'التاريخ', st.txn_type as 'النوع',
                       st.reference_no as 'المرجع',
                       i.name_ar as 'الصنف', w.name_ar as 'المستودع',
                       st.quantity as 'الكمية'
                FROM stock_transactions st
                LEFT JOIN items i ON st.item_id = i.id
                LEFT JOIN warehouses w ON st.warehouse_id = w.id
                ORDER BY st.txn_date DESC, st.id DESC
                LIMIT 1000
            """,
            'columns': ['التاريخ', 'النوع', 'المرجع', 'الصنف', 'المستودع', 'الكمية'],
        },
        'payment_allocations': {
            'name': 'توزيعات المدفوعات',
            'sql': """
                SELECT pa.payment_type as 'النوع', pa.invoice_type as 'نوع الفاتورة',
                       pa.invoice_no as 'رقم الفاتورة', pa.amount as 'المبلغ'
                FROM payment_allocations pa
                ORDER BY pa.id DESC
                LIMIT 1000
            """,
            'columns': ['النوع', 'نوع الفاتورة', 'رقم الفاتورة', 'المبلغ'],
        },
    }
    
    # =========================================
    # قوالب الاستيراد
    # =========================================
    IMPORT_TEMPLATES = {
        'accounts': {
            'name': 'الحسابات',
            'table': 'accounts',
            'columns': {
                'code': 'الكود',
                'name_ar': 'الاسم عربي',
                'name_en': 'الاسم إنجليزي',
                'level': 'المستوى',
                'nature': 'الطبيعة',
                'parent_code': 'الحساب الأب',
                'financial_statement': 'القائمة',
            },
            'required': ['code', 'name_ar'],
        },
        'customers': {
            'name': 'العملاء',
            'table': 'customers',
            'columns': {
                'code': 'الكود',
                'name_ar': 'الاسم',
                'name_en': 'الاسم إنجليزي',
                'phone': 'الهاتف',
                'email': 'البريد',
                'tax_number': 'الرقم الضريبي',
                'credit_limit': 'حد الائتمان',
                'payment_terms': 'مدة السداد',
            },
            'required': ['code', 'name_ar'],
        },
        'suppliers': {
            'name': 'الموردين',
            'table': 'suppliers',
            'columns': {
                'code': 'الكود',
                'name_ar': 'الاسم',
                'name_en': 'الاسم إنجليزي',
                'phone': 'الهاتف',
                'email': 'البريد',
                'tax_number': 'الرقم الضريبي',
                'credit_limit': 'حد الائتمان',
            },
            'required': ['code', 'name_ar'],
        },
    }
    
    # =========================================
    # التصدير
    # =========================================
    
    def export_to_excel(self, template_key, filters=None):
        """تصدير بيانات لـ Excel"""
        if template_key not in self.EXPORT_TEMPLATES:
            return None, 'قالب غير موجود'
        
        template = self.EXPORT_TEMPLATES[template_key]
        
        # جلب البيانات
        try:
            rows = db.query(template['sql'])
        except Exception as e:
            return None, str(e)
        
        if not rows:
            return None, 'لا توجد بيانات'
        
        # إنشاء Excel
        wb = Workbook()
        ws = wb.active
        ws.title = template['name']
        ws.sheet_view.rightToLeft = True
        
        # الألوان
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, size=12)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # العناوين
        columns = list(rows[0].keys())
        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        # البيانات
        for row_idx, row in enumerate(rows, 2):
            for col_idx, col_name in enumerate(columns, 1):
                value = row.get(col_name)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border
                cell.alignment = Alignment(horizontal='right', vertical='center')
        
        # عرض الأعمدة
        for col_idx, col_name in enumerate(columns, 1):
            max_length = len(str(col_name))
            for row in rows[:100]:
                val = str(row.get(col_name, ''))
                if len(val) > max_length:
                    max_length = len(val)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 4, 50)
        
        # حفظ في الذاكرة
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output, None
    
    # =========================================
    # الاستيراد
    # =========================================
    
    def import_from_excel(self, template_key, file_content):
        """استيراد من Excel"""
        if template_key not in self.IMPORT_TEMPLATES:
            return {'success': False, 'message': 'قالب غير موجود'}
        
        template = self.IMPORT_TEMPLATES[template_key]
        
        try:
            wb = load_workbook(io.BytesIO(file_content), data_only=True)
            ws = wb.active
        except Exception as e:
            return {'success': False, 'message': f'ملف Excel غير صالح: {e}'}
        
        # قراءة العناوين
        headers = []
        for col in ws[1]:
            headers.append(str(col.value).strip() if col.value else '')
        
        # خريطة الأعمدة
        col_map = {}
        for field, label in template['columns'].items():
            if label in headers:
                col_map[field] = headers.index(label)
        
        # التحقق من الأعمدة المطلوبة
        missing = []
        for field in template['required']:
            if field not in col_map:
                missing.append(template['columns'].get(field, field))
        
        if missing:
            return {
                'success': False,
                'message': f'أعمدة مطلوبة ناقصة: {", ".join(missing)}'
            }
        
        # قراءة الصفوف
        imported = 0
        updated = 0
        errors = []
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue
            
            data = {}
            for field, col_idx in col_map.items():
                value = row[col_idx] if col_idx < len(row) else None
                data[field] = value
            
            # تحقق من الحقول المطلوبة
            if not all(data.get(f) for f in template['required']):
                continue
            
            # إدخال أو تحديث
            try:
                table = template['table']
                
                # تحقق إن موجود
                existing = db.query(
                    f"SELECT id FROM {table} WHERE code = ?",
                    (data['code'],),
                    one=True
                )
                
                if existing:
                    # تحديث
                    fields = [f"{k} = ?" for k in data.keys() if k != 'code']
                    values = [data[k] for k in data.keys() if k != 'code']
                    values.append(data['code'])
                    
                    db.execute(
                        f"UPDATE {table} SET {', '.join(fields)} WHERE code = ?",
                        values
                    )
                    updated += 1
                else:
                    # إدخال
                    cols = ', '.join(data.keys())
                    placeholders = ', '.join(['?' for _ in data])
                    values = list(data.values())
                    
                    db.execute(
                        f"INSERT INTO {table} ({cols}) VALUES ({placeholders})",
                        values
                    )
                    imported += 1
            except Exception as e:
                errors.append(f'صف {row_idx}: {str(e)[:80]}')
        
        return {
            'success': True,
            'data': {
                'imported': imported,
                'updated': updated,
                'total': imported + updated,
                'errors': errors[:10]
            }
        }
    
    # =========================================
    # قالب فارغ للاستيراد
    # =========================================
    
    def get_import_template(self, template_key):
        """إنشاء ملف Excel فاضي للاستيراد"""
        if template_key not in self.IMPORT_TEMPLATES:
            return None
        
        template = self.IMPORT_TEMPLATES[template_key]
        
        wb = Workbook()
        ws = wb.active
        ws.title = template['name']
        ws.sheet_view.rightToLeft = True
        
        # الألوان
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        required_fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, size=12)
        
        # العناوين
        for col_idx, (field, label) in enumerate(template['columns'].items(), 1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            if field in template['required']:
                cell.fill = required_fill
                cell.value = f"{label} *"
            else:
                cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            
            ws.column_dimensions[get_column_letter(col_idx)].width = 20
        
        # صف فاضي
        for col_idx in range(1, len(template['columns']) + 1):
            ws.cell(row=2, column=col_idx, value='')
        
        # حفظ في الذاكرة
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output
    
    def list_templates(self):
        """قائمة القوالب المتاحة"""
        return {
            'export': [
                {'key': k, 'name': v['name']}
                for k, v in self.EXPORT_TEMPLATES.items()
            ],
            'import': [
                {'key': k, 'name': v['name']}
                for k, v in self.IMPORT_TEMPLATES.items()
            ]
        }


excel_service = ExcelService()