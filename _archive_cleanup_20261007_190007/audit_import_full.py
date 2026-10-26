"""
مراجعة شاملة لقوالب الاستيراد
- فحص كل قالب
- اختبار فعلي لاستيراد ملف تجريبي
- كشف الأخطاء
"""
import os
import io
import sqlite3
import re

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE, 'database', 'warehouse.db')
SERVICE_PATH = os.path.join(BASE, 'app', 'services', 'excel_service.py')

print("=" * 70)
print("🔍 مراجعة شاملة لقوالب الاستيراد")
print("=" * 70)

# ============================================
# 1. اتصال الداتابيز
# ============================================
conn = sqlite3.connect(DB_PATH)
conn.row_factory = sqlite3.Row

# جدول الأعمدة
tables_cols = {}
cur = conn.execute("SELECT name FROM sqlite_master WHERE type='table'")
for r in cur.fetchall():
    t = r['name']
    cur2 = conn.execute(f"PRAGMA table_info({t})")
    tables_cols[t] = [c[1] for c in cur2.fetchall()]

print(f"\n📊 عدد الجداول: {len(tables_cols)}")

# ============================================
# 2. قراءة excel_service.py
# ============================================
with open(SERVICE_PATH, 'r', encoding='utf-8') as f:
    content = f.read()

# استخراج IMPORT_TEMPLATES
import_match = re.search(r"IMPORT_TEMPLATES = \{(.*?)\n    \}", content, re.DOTALL)
if not import_match:
    print("❌ ما لقيتش IMPORT_TEMPLATES")
    exit(1)

import_block = import_match.group(1)

# استخرج كل قالب
import_pattern = r"'(?P<key>\w+)':\s*\{\s*'name':\s*'(?P<name>[^']+)',\s*'table':\s*'(?P<table>[^']+)',\s*'columns':\s*\{(?P<cols>.*?)\},\s*'required':\s*\[(?P<req>[^\]]+)\]"
import_templates = re.findall(import_pattern, import_block, re.DOTALL)

print(f"\n📊 عدد قوالب الاستيراد: {len(import_templates)}")

# ============================================
# 3. فحص كل قالب
# ============================================
print("\n" + "=" * 70)
print("📋 فحص تفصيلي")
print("=" * 70)

results = []

for key, name, table, cols_str, req_str in import_templates:
    print(f"\n🔍 {key} — {name}")
    
    # فحص 1: الجدول موجود؟
    if table not in tables_cols:
        print(f"  ❌ الجدول {table} مش موجود")
        results.append((key, 'FAIL', 'table not found'))
        continue
    
    actual_cols = tables_cols[table]
    print(f"  ✅ الجدول موجود: {table}")
    
    # فحص 2: الأعمدة
    cols = re.findall(r"'(?P<field>\w+)':\s*'(?P<label>[^']+)'", cols_str)
    fields = [c[0] for c in cols]
    labels = [c[1] for c in cols]
    
    missing = [f for f in fields if f not in actual_cols and f != 'id']
    
    if missing:
        print(f"  ⚠️  أعمدة ناقصة: {', '.join(missing)}")
    else:
        print(f"  ✅ كل الأعمدة موجودة ({len(fields)})")
    
    # فحص 3: الأعمدة المطلوبة
    req_fields = [r.strip().strip("'").strip() for r in req_str.split(',')]
    req_missing = [f for f in req_fields if f not in fields]
    
    if req_missing:
        print(f"  ⚠️  أعمدة مطلوبة ناقصة: {req_missing}")
    else:
        print(f"  ✅ الأعمدة المطلوبة كاملة: {req_fields}")
    
    # فحص 4: هل فيه unique constraint?
    has_code = 'code' in actual_cols
    if has_code:
        print(f"  ✅ يدعم INSERT/UPDATE (فيه code)")
    else:
        print(f"  ⚠️  مفيش code (INSERT فقط)")
    
    # فحص 5: تحقق من mapping
    mapping_match = re.search(rf"'{key}':\s*\{{.*?'mapping':\s*\{{(.*?)\}}", import_block, re.DOTALL)
    if mapping_match:
        print(f"  ⚠️  يحتوي على mapping (يحتاج فحص إضافي)")
    
    # النتيجة
    if missing:
        results.append((key, 'WARN', f'missing: {missing}'))
    else:
        results.append((key, 'OK', 'all good'))

# ============================================
# 4. اختبار فعلي — إنشاء ملف Excel واستيراده
# ============================================
print("\n" + "=" * 70)
print("🧪 اختبار فعلي")
print("=" * 70)

try:
    from openpyxl import Workbook
    
    for key, name, table, cols_str, req_str in import_templates[:5]:  # أول 5 بس
        print(f"\n🔬 اختبار: {key}")
        
        # استخرج الأعمدة
        cols = re.findall(r"'(?P<field>\w+)':\s*'(?P<label>[^']+)'", cols_str)
        
        # إنشاء ملف Excel تجريبي
        wb = Workbook()
        ws = wb.active
        
        # العناوين
        for col_idx, (field, label) in enumerate(cols, 1):
            ws.cell(row=1, column=col_idx, value=label)
        
        # صف واحد تجريبي
        test_code = f"TEST-{key.upper()}-001"
        for col_idx, (field, label) in enumerate(cols, 1):
            if field == 'code':
                ws.cell(row=2, column=col_idx, value=test_code)
            elif field == 'name_ar':
                ws.cell(row=2, column=col_idx, value=f'عنصر تجريبي - {name}')
            elif field == 'name_en':
                ws.cell(row=2, column=col_idx, value=f'Test - {name}')
            elif field == 'level':
                ws.cell(row=2, column=col_idx, value='L4')
            elif field == 'nature':
                ws.cell(row=2, column=col_idx, value='D')
            else:
                ws.cell(row=2, column=col_idx, value=f'test_{field}')
        
        # حفظ في الذاكرة
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        file_content = output.getvalue()
        
        print(f"  ✅ تم إنشاء ملف Excel ({len(file_content)} bytes)")
        
        # اختبار الاستيراد (dry run)
        try:
            from app.services.excel_service import excel_service
            result = excel_service.import_from_excel(key, file_content, dry_run=True)
            
            if result.get('success'):
                d = result.get('data', {})
                print(f"  ✅ Dry run نجح: {d.get('total_rows', 0)} صف")
                if d.get('errors'):
                    print(f"  ⚠️  {len(d['errors'])} خطأ")
                    for e in d['errors'][:2]:
                        print(f"     • {e}")
            else:
                print(f"  ❌ فشل: {result.get('message')}")
        except Exception as e:
            print(f"  ❌ خطأ في الاستيراد: {str(e)[:80]}")

except ImportError as e:
    print(f"⚠️  openpyxl مش متثبت: {e}")

# ============================================
# 5. التقرير النهائي
# ============================================
print("\n" + "=" * 70)
print("📊 التقرير النهائي")
print("=" * 70)

ok = sum(1 for _, s, _ in results if s == 'OK')
warn = sum(1 for _, s, _ in results if s == 'WARN')
fail = sum(1 for _, s, _ in results if s == 'FAIL')

print(f"\n  ✅ سليم: {ok}")
print(f"  ⚠️  تحذير: {warn}")
print(f"  ❌ فشل: {fail}")

if warn or fail:
    print(f"\n  📋 القوالب المحتاجة مراجعة:")
    for key, status, msg in results:
        if status != 'OK':
            print(f"     • {key}: {msg}")

# حفظ التقرير
report_path = os.path.join(BASE, 'import_report.txt')
with open(report_path, 'w', encoding='utf-8') as f:
    f.write("تقرير مراجعة قوالب الاستيراد\n")
    f.write("=" * 60 + "\n\n")
    f.write(f"إجمالي: {len(results)}\n")
    f.write(f"✅ سليم: {ok}\n")
    f.write(f"⚠️  تحذير: {warn}\n")
    f.write(f"❌ فشل: {fail}\n\n")
    
    f.write("التفاصيل:\n")
    for key, status, msg in results:
        f.write(f"  [{status}] {key}: {msg}\n")

print(f"\n📄 التقرير محفوظ: {report_path}")

conn.close()